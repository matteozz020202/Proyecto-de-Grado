"""Importación controlada del Core sintético v1.1, sin modificar el esquema."""
from collections import Counter, defaultdict
from datetime import date, datetime, time
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from django.contrib.auth.models import User
from django.core.exceptions import ValidationError
from django.db import connection, transaction
from django.utils import timezone
from openpyxl import load_workbook

from courts.models import Court, CourtSchedule, Venue
from .models import Payment, Reservation


class DatasetError(ValueError):
    pass


SHEETS = {
    "Users": "user_id", "Venues": "venue_id", "Courts": "court_id",
    "CourtSchedule": "schedule_id", "Reservations": "reservation_id",
    "Payments": "payment_id",
}


def read_core(path):
    path = Path(path)
    if path.suffix.lower() == ".zip":
        with ZipFile(path) as archive:
            candidates = [name for name in archive.namelist()
                          if Path(name).name == "Cancha_Lista_Core_v1_1_Ajustado.xlsx"]
            if len(candidates) != 1:
                raise DatasetError("El ZIP debe contener exactamente un Core v1.1 ajustado.")
            source = BytesIO(archive.read(candidates[0]))
    else:
        source = path
    workbook = load_workbook(source, read_only=True, data_only=True)
    data = {}
    try:
        for name, id_field in SHEETS.items():
            if name not in workbook.sheetnames:
                raise DatasetError(f"Falta la hoja {name}.")
            rows = iter(workbook[name].values)
            headers = next(rows)
            if id_field not in headers or len(headers) != len(set(headers)):
                raise DatasetError(f"Encabezados inválidos en {name}.")
            data[name] = [dict(zip(headers, row)) for row in rows
                          if any(value is not None for value in row)]
            ids = [row[id_field] for row in data[name]]
            if any(type(value) is not int or value <= 0 for value in ids):
                raise DatasetError(f"IDs inválidos en {name}.")
            if len(ids) != len(set(ids)):
                raise DatasetError(f"IDs duplicados en {name}.")
    finally:
        workbook.close()
    return data


def aware(value):
    if value is None:
        return None
    if not isinstance(value, datetime):
        raise DatasetError("Se esperaba una fecha/hora Excel.")
    return timezone.make_aware(value) if timezone.is_naive(value) else value


def money(value):
    result = Decimal(str(value))
    if not result.is_finite() or result < 0 or result != result.quantize(Decimal(".01")):
        raise DatasetError(f"Monto inválido: {value}.")
    return result


def validate_overlap(reservations):
    groups = defaultdict(list)
    for reservation in reservations:
        if reservation.reservation_status not in ("CANCELLED", "EXPIRED"):
            groups[(reservation.court_id, reservation.reservation_date)].append(reservation)
    for group in groups.values():
        ordered = sorted(group, key=lambda r: r.start_time)
        end = None
        for reservation in ordered:
            if end is not None and reservation.start_time < end:
                raise DatasetError(f"Solapamiento en reserva {reservation.pk}.")
            end = reservation.end_time


def build_objects(data):
    """Valida sin consultas ORM; los IDs de User se mapean durante la carga."""
    for source, field, target, target_field in [
        ("Courts", "venue_id", "Venues", "venue_id"),
        ("CourtSchedule", "court_id", "Courts", "court_id"),
        ("Reservations", "court_id", "Courts", "court_id"),
        ("Reservations", "user_id", "Users", "user_id"),
        ("Payments", "reservation_id", "Reservations", "reservation_id"),
    ]:
        ids = {row[target_field] for row in data[target]}
        if any(row[field] not in ids for row in data[source]):
            raise DatasetError(f"Referencia inexistente: {source}.{field}.")

    objects = {model: [] for model in (Venue, Court, CourtSchedule, Reservation, Payment)}
    for row in data["Users"]:
        if not isinstance(row["full_name"], str) or len(row["full_name"]) > 150:
            raise DatasetError("Nombre de usuario sintético inválido.")
        if not isinstance(row["email"], str) or len(row["email"]) > 254:
            raise DatasetError("Email sintético inválido.")
        aware(row["created_at"])
    for row in data["Venues"]:
        fraction = Decimal(str(row["deposit_percentage"]))
        if not 0 <= fraction <= 1:
            raise DatasetError("El porcentaje del Core debe estar entre 0 y 1.")
        objects[Venue].append(Venue(
            venue_id=row["venue_id"], venue_name=row["venue_name"], city=row["city"],
            address=row["address"], deposit_percentage=fraction * 100,
            is_active=row["is_active"],
        ))
    for row in data["Courts"]:
        objects[Court].append(Court(
            court_id=row["court_id"], venue_id=row["venue_id"],
            court_name=row["court_name"], price_per_hour=money(row["price_per_hour"]),
            is_active=row["is_active"],
        ))
    for row in data["CourtSchedule"]:
        if row["day_of_week"] not in range(1, 8) or row["slot_duration_minutes"] <= 0:
            raise DatasetError("Horario inválido: se requieren días 1–7 y duración positiva.")
        if not isinstance(row["start_time"], time) or not isinstance(row["end_time"], time):
            raise DatasetError("Horas inválidas.")
        if row["start_time"] >= row["end_time"]:
            raise DatasetError("Horario con inicio posterior al fin.")
        objects[CourtSchedule].append(CourtSchedule(**row))
    for row in data["Reservations"]:
        if row["data_source"] != "SYNTHETIC":
            raise DatasetError("Solo se admite data_source=SYNTHETIC.")
        values = dict(row)
        value = values["reservation_date"]
        if not isinstance(value, date):
            raise DatasetError("Fecha de reserva inválida.")
        values["reservation_date"] = value.date() if isinstance(value, datetime) else value
        for field in ("created_at", "hold_expires_at", "cancelled_at"):
            values[field] = aware(values[field])
        for field in ("total_amount", "deposit_required", "remaining_amount"):
            values[field] = money(values[field])
        if values["total_amount"] != values["deposit_required"] + values["remaining_amount"]:
            raise DatasetError("Total, abono y saldo inconsistentes.")
        if not isinstance(values["start_time"], time) or not isinstance(values["end_time"], time):
            raise DatasetError("Horas de reserva inválidas.")
        if values["start_time"] >= values["end_time"]:
            raise DatasetError("Reserva con inicio posterior al fin.")
        if values["reservation_status"] == "CANCELLED" and not values["cancelled_at"]:
            raise DatasetError("Cancelación sin fecha.")
        if values["reservation_status"] == "PENDING_PAYMENT" and not values["hold_expires_at"]:
            raise DatasetError("Reserva pendiente sin vencimiento.")
        objects[Reservation].append(Reservation(**values))
    references = set()
    approved = defaultdict(Decimal)
    for row in data["Payments"]:
        if row["transaction_reference"] in references:
            raise DatasetError("Referencia de pago duplicada.")
        references.add(row["transaction_reference"])
        values = dict(row)
        values["amount"] = money(values["amount"])
        for field in ("created_at", "approved_at"):
            values[field] = aware(values[field])
        if values["payment_status"] == "APPROVED":
            if not values["approved_at"]:
                raise DatasetError("Pago aprobado sin fecha.")
            approved[values["reservation_id"]] += values["amount"]
        objects[Payment].append(Payment(**values))
    for reservation in objects[Reservation]:
        paid = approved[reservation.pk]
        if reservation.reservation_status in ("COMPLETED", "CONFIRMED", "CANCELLED"):
            if paid < reservation.deposit_required:
                raise DatasetError(f"Reserva {reservation.pk} formalizada sin abono aprobado.")
        elif paid:
            raise DatasetError(f"Reserva {reservation.pk} pendiente/expirada con pago aprobado.")
    # Field.clean valida tipos, choices, longitudes y decimales sin consultar FKs.
    for model, entries in objects.items():
        for entry in entries:
            for field in model._meta.concrete_fields:
                if not field.is_relation:
                    try:
                        field.clean(getattr(entry, field.attname), entry)
                    except ValidationError as exc:
                        raise DatasetError(f"{model.__name__} {entry.pk}: {exc}") from exc
    validate_overlap(objects[Reservation])
    return objects


def summarize(data):
    return {
        "filas": {name: len(rows) for name, rows in data.items()},
        "estados": dict(Counter(row["reservation_status"] for row in data["Reservations"])),
        "fecha_min": min(row["reservation_date"].isoformat() for row in data["Reservations"]),
        "fecha_max": max(row["reservation_date"].isoformat() for row in data["Reservations"]),
    }


@transaction.atomic
def import_core(data, objects, namespace="v1_1"):
    if connection.vendor == "postgresql":
        # Serializa importaciones y evita escrituras concurrentes en estas tablas.
        tables = [User._meta.db_table] + [model._meta.db_table for model in objects]
        with connection.cursor() as cursor:
            cursor.execute("LOCK TABLE " + ", ".join(connection.ops.quote_name(t) for t in tables)
                           + " IN SHARE ROW EXCLUSIVE MODE")
    counts = {}
    usernames = {row["user_id"]: f"synthetic_{namespace}_{row['user_id']:04d}" for row in data["Users"]}
    existing_users = {user.username: user for user in User.objects.filter(username__in=usernames.values())}
    new_users = []
    for row in data["Users"]:
        username = usernames[row["user_id"]]
        user = existing_users.get(username)
        if user:
            if user.has_usable_password() or user.is_staff or user.is_superuser or user.is_active:
                raise DatasetError(f"Colisión con una cuenta existente: {username}.")
            if user.email != row["email"] or user.first_name != row["full_name"]:
                raise DatasetError(f"Usuario sintético diferente: {username}.")
        else:
            user = User(username=username, email=row["email"], first_name=row["full_name"],
                        is_active=False, date_joined=aware(row["created_at"]))
            user.set_unusable_password()
            new_users.append(user)
    User.objects.bulk_create(new_users, batch_size=500)
    user_map = dict(User.objects.filter(username__in=usernames.values()).values_list("username", "pk"))
    counts["Users"] = len(new_users)
    for reservation in objects[Reservation]:
        reservation.user_id = user_map[usernames[reservation.user_id]]
    incoming_ids = {r.pk for r in objects[Reservation]}
    existing_other = list(Reservation.objects.filter(
        court_id__in=[c.pk for c in objects[Court]],
        reservation_date__range=(min(r.reservation_date for r in objects[Reservation]),
                                 max(r.reservation_date for r in objects[Reservation])),
    ).exclude(pk__in=incoming_ids))
    validate_overlap(existing_other + objects[Reservation])
    for model, entries in objects.items():
        existing = model.objects.in_bulk([entry.pk for entry in entries])
        new = []
        original_times = {}
        for entry in entries:
            previous = existing.get(entry.pk)
            if previous:
                for field in model._meta.concrete_fields:
                    if field.name in ("image", "description"):
                        continue
                    old, proposed = getattr(previous, field.attname), getattr(entry, field.attname)
                    # No revierte una expiración lazy ocurrida después de la carga.
                    if model is Reservation and field.name == "reservation_status":
                        if proposed == "PENDING_PAYMENT" and old == "EXPIRED":
                            continue
                    if old != proposed:
                        raise DatasetError(f"Colisión: {model.__name__} {entry.pk}, campo {field.name}.")
            else:
                new.append(entry)
                if hasattr(entry, "created_at"):
                    original_times[entry.pk] = entry.created_at
        if model is Payment:
            refs = [p.transaction_reference for p in new]
            if Payment.objects.filter(transaction_reference__in=refs).exists():
                raise DatasetError("Una referencia de pago ya pertenece a otro registro.")
        model.objects.bulk_create(new, batch_size=500)
        if original_times:
            for entry in new:
                entry.created_at = original_times[entry.pk]
            model.objects.bulk_update(new, ["created_at"], batch_size=500)
        counts[model.__name__] = len(new)
    # IDs explícitos requieren avanzar las secuencias, sin bajarlas ni reutilizar IDs.
    if connection.vendor == "postgresql":
        with connection.cursor() as cursor:
            for model in objects:
                table, column = model._meta.db_table, model._meta.pk.column
                cursor.execute("SELECT pg_get_serial_sequence(%s, %s)", [table, column])
                sequence = cursor.fetchone()[0]
                if sequence:
                    quoted_sequence = ".".join(connection.ops.quote_name(part) for part in sequence.split("."))
                    cursor.execute("SELECT last_value FROM " + quoted_sequence)
                    last_value = cursor.fetchone()[0]
                    maximum = model.objects.order_by("-pk").values_list("pk", flat=True).first() or 1
                    cursor.execute("SELECT setval(%s, %s, true)", [sequence, max(last_value, maximum)])
    return counts
